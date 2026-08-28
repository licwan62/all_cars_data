from __future__ import annotations

import argparse
import csv
import hashlib
import json
import math
import re
from collections import Counter, defaultdict
from copy import deepcopy
from dataclasses import dataclass, replace
from datetime import datetime
from pathlib import Path
from typing import Any, Iterable

from openpyxl import Workbook
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo


ROOT = Path(__file__).resolve().parents[1]
SIZE_SOURCE = ROOT / "source" / "尺码分析.csv"
RULE_SOURCE = ROOT / "尺码计算" / "input" / "尺码匹配规则.csv"
OUTPUT_DIR = ROOT / "尺码簇分析" / "output"
ANALYSIS_BOOK = OUTPUT_DIR / "尺码上限微调分析.xlsx"
MANUAL_BOOK = OUTPUT_DIR / "边界负面影响及手动调整.xlsx"
REPORT = OUTPUT_DIR / "尺码分组微调分析报告.md"
ADJUSTMENT_CSV = OUTPUT_DIR / "分组调整表.csv"

BASE_ROUGH_LIMIT = {"两厢车": 120, "跑车": 120, "三厢车": 120, "越野车": 150, "皮卡": 250}
PASSENGER_CLASSES = {"两厢车", "跑车", "三厢车", "越野车"}
HIGH_FAMILY_SALES = 250_000
VERY_HIGH_ROW_SALES = 100_000
MAX_SHIFT_DEFAULT = 100
MAX_SHIFT_PICKUP = 150
BOUNDARY_BAND_MM = 25

NAVY = "17324D"
TEAL = "0F766E"
BLUE = "2563EB"
ORANGE = "D97706"
RED = "B91C1C"
GREEN = "15803D"
LIGHT_BLUE = "DBEAFE"
LIGHT_TEAL = "CCFBF1"
LIGHT_ORANGE = "FEF3C7"
LIGHT_RED = "FEE2E2"
LIGHT_GRAY = "F1F5F9"
MID_GRAY = "CBD5E1"
WHITE = "FFFFFF"


def clean_text(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def clean_number(value: Any) -> float | None:
    if value is None or value == "":
        return None
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def normalize_version(value: Any) -> str:
    text = clean_text(value)
    return "DRW" if "DRW" in text.upper() else text


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def excel_serialized(value: Any) -> Any:
    if isinstance(value, datetime):
        return value.isoformat()
    return value


@dataclass(frozen=True)
class SizeRule:
    row_no: int
    logical_size: str
    internal_size: str
    sequence: float
    category: str
    cab: str
    version: str
    length_upper: float
    rough_upper: float
    use: str

    @property
    def pool_key(self) -> tuple[str, str, str]:
        return self.category, self.cab, self.version

    @property
    def lane(self) -> str:
        base = BASE_ROUGH_LIMIT.get(self.category)
        if self.version == "DRW":
            return "DRW"
        if base is not None and self.rough_upper > base:
            return "W"
        return "普通"


def read_source(
    rule_path: Path, size_path: Path
) -> tuple[list[SizeRule], list[dict[str, Any]], list[list[Any]]]:
    """读取项目规则和 source 下的统一尺码分析结果。"""
    rule_columns = [
        "逻辑尺码",
        "内部尺码",
        "档位序号",
        "分类",
        "CAB",
        "版本",
        "长上限",
        "参考插片上限",
        "使用",
    ]
    rules: list[SizeRule] = []
    with rule_path.open("r", encoding="utf-8-sig", newline="") as handle:
        reader = csv.DictReader(handle)
        headers = reader.fieldnames or []
        missing = [column for column in rule_columns if column not in headers]
        if missing:
            raise ValueError(f"{rule_path} 缺少字段: {', '.join(missing)}")
        rule_snapshot: list[list[Any]] = [headers]
        for row_no, row in enumerate(reader, start=2):
            rule_snapshot.append([row.get(header, "") for header in headers])
            logical = row.get("逻辑尺码")
            internal = row.get("内部尺码")
            sequence = row.get("档位序号")
            category = row.get("分类")
            cab = row.get("CAB")
            version = row.get("版本")
            length_upper = row.get("长上限")
            rough_upper = row.get("参考插片上限")
            use = row.get("使用")
            if clean_text(use).lower() != "y" or not clean_text(category):
                continue
            if (
                clean_number(sequence) is None
                or clean_number(length_upper) is None
                or clean_number(rough_upper) is None
            ):
                continue
            rules.append(
                SizeRule(
                    row_no=row_no,
                    logical_size=clean_text(logical),
                    internal_size=clean_text(internal),
                    sequence=float(sequence),
                    category=clean_text(category),
                    cab=clean_text(cab),
                    version=normalize_version(version),
                    length_upper=float(length_upper),
                    rough_upper=float(rough_upper),
                    use=clean_text(use),
                )
            )

    required_size_columns = [
        "DIMENSION-ID",
        "自动尺码",
        "MAKE",
        "MODEL",
        "版本",
        "结构",
        "CAB",
        "YEAR",
        "分类",
        "L-MM",
        "W-MM",
        "H-MM",
        "销量合计",
        "参考插片",
    ]
    records: list[dict[str, Any]] = []
    with size_path.open("r", encoding="utf-8-sig", newline="") as handle:
        reader = csv.DictReader(handle)
        headers = reader.fieldnames or []
        missing = [column for column in required_size_columns if column not in headers]
        if missing:
            raise ValueError(f"{size_path} 缺少字段: {', '.join(missing)}")
        for source_row, record in enumerate(reader, start=2):
            if not clean_text(record.get("DIMENSION-ID")):
                continue
            record["__source_row"] = source_row
            record["__sales"] = max(clean_number(record.get("销量合计")) or 0.0, 0.0)
            record["__length"] = clean_number(record.get("L-MM"))
            record["__width"] = clean_number(record.get("W-MM"))
            record["__height"] = clean_number(record.get("H-MM"))
            record["__rough"] = clean_number(record.get("参考插片"))
            if (
                record["__rough"] is None
                and record["__width"] is not None
                and record["__height"] is not None
            ):
                record["__rough"] = record["__width"] + 2 * record["__height"]
            record["__category"] = clean_text(record.get("分类"))
            record["__cab"] = clean_text(record.get("CAB"))
            record["__version"] = normalize_version(record.get("版本"))
            records.append(record)

    if not rules:
        raise ValueError(f"{rule_path} 没有启用的有效规则")
    if not records:
        raise ValueError(f"{size_path} 没有有效 DIMENSION-ID 数据")
    return rules, records, rule_snapshot


def build_pools(rules: Iterable[SizeRule]) -> dict[tuple[str, str, str], list[SizeRule]]:
    pools: dict[tuple[str, str, str], list[SizeRule]] = defaultdict(list)
    for rule in rules:
        pools[rule.pool_key].append(rule)
    for pool in pools.values():
        pool.sort(key=lambda r: (r.sequence, r.length_upper))
    return pools


def try_pool(
    pool: list[SizeRule] | None,
    length: float,
    rough: float,
    length_tolerance: float = 500,
) -> SizeRule | None:
    if not pool:
        return None
    for rule in pool:
        if rule.length_upper >= length and rule.rough_upper >= rough:
            if rule.length_upper - length <= length_tolerance:
                return rule
            return None
    return None


def match_rule(
    record: dict[str, Any], pools: dict[tuple[str, str, str], list[SizeRule]]
) -> SizeRule | None:
    length = record["__length"]
    width = record["__width"]
    height = record["__height"]
    rough = record["__rough"]
    category = record["__category"]
    cab = record["__cab"]
    version = record["__version"]
    if length is None or width is None or height is None or rough is None or not category:
        return None

    def match_category(target_category: str) -> SizeRule | None:
        keys: list[tuple[str, str, str]] = []
        if cab and version:
            keys.append((target_category, cab, version))
        if version:
            keys.append((target_category, "", version))
        keys.append((target_category, "", ""))
        seen: set[tuple[str, str, str]] = set()
        for key in keys:
            if key in seen:
                continue
            seen.add(key)
            candidate = try_pool(pools.get(key), length, rough)
            if candidate is not None:
                return candidate
        return None

    if category == "三厢车":
        max_lengths = []
        for key in ((category, cab, version), (category, "", version), (category, "", "")):
            if key in pools and pools[key]:
                max_lengths.append(max(rule.length_upper for rule in pools[key]))
        if max_lengths and length > max(max_lengths):
            return match_category("跑车")
    return match_category(category)


def assign_records(records: list[dict[str, Any]], rules: list[SizeRule]) -> list[dict[str, Any]]:
    assigned: list[dict[str, Any]] = []
    pools = build_pools(rules)
    for record in records:
        rule = match_rule(record, pools)
        new_record = dict(record)
        new_record["__matched_rule"] = rule
        new_record["__logical_size"] = rule.logical_size if rule else ""
        new_record["__internal_size"] = rule.internal_size if rule else ""
        new_record["__length_surplus"] = (
            rule.length_upper - record["__length"]
            if rule and record["__length"] is not None
            else None
        )
        new_record["__lane"] = rule.lane if rule else lane_for_record(record)
        assigned.append(new_record)
    return assigned


def lane_for_record(record: dict[str, Any]) -> str:
    if record.get("__version") == "DRW":
        return "DRW"
    base = BASE_ROUGH_LIMIT.get(record.get("__category"))
    rough = record.get("__rough")
    if base is not None and rough is not None and rough > base:
        return "W"
    return "普通"


def family_key(record: dict[str, Any]) -> tuple[str, ...]:
    category = record["__category"]
    base = [
        clean_text(record.get("MAKE")),
        clean_text(record.get("MODEL")),
        clean_text(record.get("结构")),
        category,
        lane_for_record(record),
    ]
    if category == "皮卡":
        base.extend(
            [
                clean_text(record.get("版本")),
                clean_text(record.get("CAB")),
                clean_text(record.get("BED")),
            ]
        )
    return tuple(base)


def family_label(key: tuple[str, ...]) -> str:
    label = " / ".join(value for value in key[:3] if value)
    if key[4] != "普通":
        label += f" [{key[4]}]"
    if len(key) > 5:
        extras = "/".join(value for value in key[5:] if value)
        if extras:
            label += f" ({extras})"
    return label


def generation_count(rows: list[dict[str, Any]]) -> int:
    return len({clean_text(r.get("代际")) or clean_text(r.get("YEAR")) for r in rows})


def weighted_modal_size(rows: list[dict[str, Any]], size_field: str) -> tuple[str, float, bool]:
    weights: dict[str, float] = defaultdict(float)
    for row in rows:
        size = clean_text(row.get(size_field))
        if size:
            weights[size] += row["__sales"]
    if not weights:
        return "", 0.0, False
    ordered = sorted(weights.items(), key=lambda item: (-item[1], item[0]))
    unique = len(ordered) == 1 or ordered[0][1] > ordered[1][1]
    return ordered[0][0], ordered[0][1], unique


def group_families(rows: list[dict[str, Any]]) -> dict[tuple[str, ...], list[dict[str, Any]]]:
    result: dict[tuple[str, ...], list[dict[str, Any]]] = defaultdict(list)
    for row in rows:
        result[family_key(row)].append(row)
    return result


def family_metrics(rows: list[dict[str, Any]], size_field: str) -> dict[tuple[str, ...], dict[str, Any]]:
    result: dict[tuple[str, ...], dict[str, Any]] = {}
    for key, family_rows in group_families(rows).items():
        total_sales = sum(r["__sales"] for r in family_rows)
        modal_size, modal_sales, unique = weighted_modal_size(family_rows, size_field)
        sizes = sorted({clean_text(r.get(size_field)) for r in family_rows if clean_text(r.get(size_field))})
        result[key] = {
            "rows": family_rows,
            "total_sales": total_sales,
            "modal_size": modal_size,
            "modal_sales": modal_sales,
            "modal_unique": unique,
            "purity": modal_sales / total_sales if total_sales > 0 else 0,
            "sizes": sizes,
            "size_sales": dict(
                (
                    size,
                    sum(r["__sales"] for r in family_rows if clean_text(r.get(size_field)) == size),
                )
                for size in sizes
            ),
            "generations": generation_count(family_rows),
            "high": total_sales >= HIGH_FAMILY_SALES and generation_count(family_rows) >= 2,
        }
    return result


def verify_current_assignments(records: list[dict[str, Any]], rules: list[SizeRule]) -> list[dict[str, Any]]:
    assigned = assign_records(records, rules)
    mismatches = []
    for row in assigned:
        sheet_size = clean_text(row.get("自动尺码"))
        calculated = row["__internal_size"]
        if sheet_size in {"数据不全", "无可用尺码"}:
            sheet_size = ""
        if sheet_size != calculated:
            mismatches.append(row)
    return mismatches


def boundary_rules(rules: list[SizeRule]) -> list[SizeRule]:
    by_pool_lane: dict[tuple[tuple[str, str, str], str], list[SizeRule]] = defaultdict(list)
    for rule in rules:
        by_pool_lane[(rule.pool_key, rule.lane)].append(rule)
    output: list[SizeRule] = []
    for lane_rules in by_pool_lane.values():
        lane_rules.sort(key=lambda r: (r.sequence, r.length_upper))
        output.extend(lane_rules[:-1])
    return sorted(output, key=lambda r: (r.category, r.version, r.lane, r.sequence, r.length_upper))


def candidate_thresholds(rule: SizeRule, rules: list[SizeRule], records: list[dict[str, Any]]) -> list[int]:
    shift = MAX_SHIFT_PICKUP if rule.category == "皮卡" else MAX_SHIFT_DEFAULT
    lower = int(math.floor(rule.length_upper - shift))
    upper = int(math.ceil(rule.length_upper + shift))
    same_lane = sorted(
        [r for r in rules if r.pool_key == rule.pool_key and r.lane == rule.lane],
        key=lambda r: (r.sequence, r.length_upper),
    )
    index = same_lane.index(rule)
    if index > 0:
        lower = max(lower, int(math.floor(same_lane[index - 1].length_upper)) + 1)
    if index + 1 < len(same_lane):
        upper = min(upper, int(math.ceil(same_lane[index + 1].length_upper)) - 1)
    values = {int(round(rule.length_upper))}
    for record in records:
        if record["__category"] != rule.category or lane_for_record(record) != rule.lane:
            continue
        length = record["__length"]
        if length is None or length < lower or length > upper:
            continue
        values.add(int(round(length)))
        values.add(int(math.floor(length)) - 1)
    return sorted(v for v in values if lower <= v <= upper)


def replace_rule(rules: list[SizeRule], target: SizeRule, new_upper: int) -> list[SizeRule]:
    return [replace(rule, length_upper=float(new_upper)) if rule.row_no == target.row_no else rule for rule in rules]


def evaluate_candidate(
    target: SizeRule,
    new_upper: int,
    rules: list[SizeRule],
    base_rows: list[dict[str, Any]],
    base_family_metrics: dict[tuple[str, ...], dict[str, Any]],
) -> dict[str, Any]:
    candidate_rules = replace_rule(rules, target, new_upper)
    candidate_pools = build_pools(candidate_rules)

    changed = []
    changed_by_family: dict[tuple[str, ...], list[tuple[dict[str, Any], dict[str, Any]]]] = defaultdict(list)
    for before in base_rows:
        if before["__category"] != target.category:
            continue
        candidate_rule = match_rule(before, candidate_pools)
        candidate_size = candidate_rule.logical_size if candidate_rule else ""
        if before["__logical_size"] == candidate_size:
            continue
        after = dict(before)
        after["__matched_rule"] = candidate_rule
        after["__logical_size"] = candidate_size
        after["__internal_size"] = candidate_rule.internal_size if candidate_rule else ""
        after["__length_surplus"] = (
            candidate_rule.length_upper - before["__length"]
            if candidate_rule and before["__length"] is not None
            else None
        )
        changed.append((before, after))
        changed_by_family[family_key(before)].append((before, after))

    improve_sales = 0.0
    worsen_sales = 0.0
    improved_families: list[tuple[str, ...]] = []
    worsened_families: list[tuple[str, ...]] = []
    candidate_metric_overrides: dict[tuple[str, ...], dict[str, Any]] = {}
    for key, family_changes in changed_by_family.items():
        before_metric = base_family_metrics[key]
        size_sales = defaultdict(float, before_metric["size_sales"])
        for before_row, after_row in family_changes:
            size_sales[before_row["__logical_size"]] -= before_row["__sales"]
            size_sales[after_row["__logical_size"]] += before_row["__sales"]
        size_sales = {size: sales for size, sales in size_sales.items() if sales > 0.5 and size}
        modal_sales = max(size_sales.values(), default=0.0)
        modal_sizes = sorted(size for size, sales in size_sales.items() if abs(sales - modal_sales) < 0.5)
        candidate_metric_overrides[key] = {
            "modal_sales": modal_sales,
            "modal_size": modal_sizes[0] if modal_sizes else "",
            "modal_unique": len(modal_sizes) == 1,
            "size_sales": size_sales,
        }

    for key, before in base_family_metrics.items():
        if not before["high"]:
            continue
        after = candidate_metric_overrides.get(key, before)
        delta = after["modal_sales"] - before["modal_sales"]
        if delta > 0.5:
            improve_sales += delta
            improved_families.append(key)
        elif delta < -0.5:
            worsen_sales += -delta
            worsened_families.append(key)

    moved_sales = sum(before["__sales"] for before, _ in changed)
    neutral_sales = max(moved_sales - improve_sales, 0.0)
    net_score = improve_sales - 1.5 * worsen_sales - 0.20 * neutral_sales
    near_boundary = [
        row
        for row in base_rows
        if row["__category"] == target.category
        and lane_for_record(row) == target.lane
        and row["__length"] is not None
        and abs(row["__length"] - new_upper) <= BOUNDARY_BAND_MM
    ]
    return {
        "rule": target,
        "new_upper": new_upper,
        "delta": new_upper - target.length_upper,
        "candidate_rules": candidate_rules,
        "candidate_metric_overrides": candidate_metric_overrides,
        "changed": changed,
        "changed_count": len(changed),
        "moved_sales": moved_sales,
        "improve_sales": improve_sales,
        "worsen_sales": worsen_sales,
        "neutral_sales": neutral_sales,
        "net_score": net_score,
        "improved_families": improved_families,
        "worsened_families": worsened_families,
        "near_boundary_count": len(near_boundary),
        "near_boundary_sales": sum(r["__sales"] for r in near_boundary),
    }


def analyze_candidates(
    rules: list[SizeRule], records: list[dict[str, Any]]
) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    base_rows = assign_records(records, rules)
    base_metrics = family_metrics(base_rows, "__logical_size")
    all_results: list[dict[str, Any]] = []
    best_results: list[dict[str, Any]] = []
    for boundary in boundary_rules(rules):
        evaluations = []
        for threshold in candidate_thresholds(boundary, rules, records):
            if threshold == int(round(boundary.length_upper)):
                continue
            result = evaluate_candidate(boundary, threshold, rules, base_rows, base_metrics)
            evaluations.append(result)
            all_results.append(result)
        if not evaluations:
            continue
        evaluations.sort(
            key=lambda r: (
                -r["net_score"],
                r["worsen_sales"],
                r["neutral_sales"],
                abs(r["delta"]),
            )
        )
        best_results.append(evaluations[0])
    best_results.sort(key=lambda r: (-r["net_score"], -r["improve_sales"], abs(r["delta"])))
    return best_results, all_results


def classify_candidate(result: dict[str, Any]) -> tuple[str, str]:
    improve = result["improve_sales"]
    worsen = result["worsen_sales"]
    neutral = result["neutral_sales"]
    moved = result["moved_sales"]
    delta = abs(result["delta"])
    if (
        result["net_score"] > 0
        and improve >= HIGH_FAMILY_SALES
        and worsen <= improve * 0.15
        and neutral <= improve * 1.5
        and delta <= MAX_SHIFT_DEFAULT
    ):
        return "建议全局微调", "跨代归组收益明显，且连带销量/负向纯度变化受控"
    if improve >= HIGH_FAMILY_SALES and moved > 0:
        return "改为手动调整", "存在跨代归组收益，但全局移动的非目标车型或边界密度过高"
    return "保持现状", "未发现足以覆盖连带影响的高销量跨代收益"


def policy_decision(result: dict[str, Any]) -> tuple[str, float, str]:
    rule = result["rule"]
    if rule.category == "皮卡" and rule.logical_size == "PK-XXL" and result["new_upper"] == 6462:
        return (
            "建议小范围试行",
            6462,
            "仅 2 条 F-250/F-350 Super Duty 高销量记录换入 PK-XXL，无非目标连带换档；向小一档，须先做样车覆盖验证。",
        )
    if rule.category == "皮卡" and rule.logical_size == "PK-L" and result["new_upper"] == 5916:
        return (
            "样车验证后可选",
            rule.length_upper,
            "+20 mm 会让 8 条、451 万销量记录从 PK-XL 换到 PK-L；收益存在，但连带销量高，暂不直接改规则。",
        )
    if rule.category in {"三厢车", "跑车"} and rule.logical_size == "3L":
        return (
            "保持规则，车型手动",
            rule.length_upper,
            "Accord 可通过手动归组实现；全局下调会让大量 4694–4736 mm 车型统一换大一档，长度余量突增。",
        )
    if rule.lane == "W":
        return "保持现状", rule.length_upper, "W 版长度边界未出现低连带、高收益的稳定微调点。"
    if result["improve_sales"] >= HIGH_FAMILY_SALES and result["changed_count"] > 0:
        return (
            "保持规则，车型手动",
            rule.length_upper,
            "模拟可改善部分跨代归组，但全局换档规模或向小一档的实物风险不适合直接落规则。",
        )
    return "保持现状", rule.length_upper, "微调收益不足，或没有实际车型发生换档。"


def rule_for_logical(
    rules: list[SizeRule], category: str, lane: str, logical_size: str
) -> SizeRule | None:
    candidates = [
        rule
        for rule in rules
        if rule.category == category
        and rule.lane == lane
        and rule.logical_size == logical_size
        and rule.cab == ""
        and rule.version == ""
    ]
    if not candidates:
        return None
    return sorted(candidates, key=lambda r: (r.sequence, r.length_upper))[0]


def lane_rules(rules: list[SizeRule], category: str, lane: str) -> list[SizeRule]:
    return sorted(
        [
            rule
            for rule in rules
            if rule.category == category
            and rule.lane == lane
            and rule.cab == ""
            and rule.version == ""
        ],
        key=lambda r: (r.sequence, r.length_upper),
    )


def manual_adjustment_candidates(
    base_rows: list[dict[str, Any]], rules: list[SizeRule]
) -> list[dict[str, Any]]:
    metrics = family_metrics(base_rows, "__logical_size")
    output: list[dict[str, Any]] = []
    for key, metric in metrics.items():
        if not metric["high"] or len(metric["sizes"]) < 2 or not metric["modal_unique"]:
            continue
        target_size = metric["modal_size"]
        category, lane = key[3], key[4]
        ordered = lane_rules(rules, category, lane)
        position = {rule.logical_size: index for index, rule in enumerate(ordered)}
        target_rule = rule_for_logical(rules, category, lane, target_size)
        if target_rule is None or target_size not in position:
            continue
        for row in metric["rows"]:
            current_size = row["__logical_size"]
            if not current_size or current_size == target_size or current_size not in position:
                continue
            if abs(position[current_size] - position[target_size]) != 1:
                continue
            current_rule: SizeRule | None = row.get("__matched_rule")
            length = row["__length"]
            rough = row["__rough"]
            if current_rule is None or length is None or rough is None or target_rule.rough_upper < rough:
                continue
            if position[target_size] > position[current_size]:
                boundary_distance = current_rule.length_upper - length
                if boundary_distance < -0.5 or boundary_distance > 75:
                    continue
                direction = "手动升一档"
                risk = "长度可覆盖；主要风险是余量增大、版型偏松"
            else:
                boundary_distance = length - target_rule.length_upper
                if boundary_distance < -0.5 or boundary_distance > 20:
                    continue
                direction = "验证后手动降一档"
                risk = "车型长度略超现行上限；必须先做样车/版型验证"
            is_accord_3xl_target = (
                clean_text(row.get("MAKE")) == "Honda"
                and clean_text(row.get("MODEL")) == "Accord"
                and category in {"三厢车", "跑车"}
                and current_size == "3L"
                and target_size == "3XL"
            )
            priority = "P0" if is_accord_3xl_target else ("P1" if row["__sales"] >= 250_000 else "P2")
            output.append(
                {
                    "优先级": priority,
                    "MAKE": clean_text(row.get("MAKE")),
                    "MODEL": clean_text(row.get("MODEL")),
                    "结构": clean_text(row.get("结构")),
                    "版本": clean_text(row.get("版本")),
                    "CAB": clean_text(row.get("CAB")),
                    "BED": clean_text(row.get("BED")),
                    "代际": clean_text(row.get("代际")),
                    "YEAR": clean_text(row.get("YEAR")),
                    "分类": category,
                    "通道": lane,
                    "L-MM": length,
                    "参考插片": rough,
                    "销量合计": row["__sales"],
                    "当前逻辑尺码": current_size,
                    "当前内部尺码": row["__internal_size"],
                    "建议逻辑尺码": target_size,
                    "建议内部尺码": target_rule.internal_size,
                    "建议方式": direction,
                    "距当前边界-mm": boundary_distance,
                    "建议后长度余量-mm": target_rule.length_upper - length,
                    "家族总销量": metric["total_sales"],
                    "当前主组占比": metric["purity"],
                    "说明": f"{family_label(key)} 的销量主组为 {target_size}；{risk}",
                    "DIMENSION-ID": clean_text(row.get("DIMENSION-ID")),
                    "源行号": row["__source_row"],
                }
            )
    output.sort(
        key=lambda row: (
            {"P0": 0, "P1": 1, "P2": 2}.get(row["优先级"], 9),
            -row["销量合计"],
            row["MAKE"],
            row["MODEL"],
            row["YEAR"],
        )
    )
    return output


def export_manual_adjustments_csv(manual_rows: list[dict[str, Any]]) -> dict[str, Any]:
    """Export every manual-review candidate in the import-ready four-column layout."""
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    fieldnames = ["dimension-id", "original_size", "suggest_size", "原因"]
    with ADJUSTMENT_CSV.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames)
        writer.writeheader()
        for row in manual_rows:
            reason = "；".join(
                [
                    f"优先级={row['优先级']}",
                    f"家族主组占比={row['当前主组占比']:.1%}",
                    f"家族总销量={row['家族总销量']:,.0f}",
                    f"建议方式={row['建议方式']}",
                    f"距当前边界={row['距当前边界-mm']:.0f}mm",
                    f"建议后长度余量={row['建议后长度余量-mm']:.0f}mm",
                    row["说明"],
                ]
            )
            writer.writerow(
                {
                    "dimension-id": row["DIMENSION-ID"],
                    "original_size": row["当前内部尺码"],
                    "suggest_size": row["建议内部尺码"],
                    "原因": reason,
                }
            )
    return {
        "csv": str(ADJUSTMENT_CSV),
        "rows": len(manual_rows),
        "priority_counts": dict(Counter(row["优先级"] for row in manual_rows)),
    }


def accord_boundary_scenario(
    base_rows: list[dict[str, Any]], rules: list[SizeRule]
) -> tuple[list[dict[str, Any]], list[dict[str, Any]], list[dict[str, Any]]]:
    base_metrics = family_metrics(base_rows, "__logical_size")
    all_changed: list[dict[str, Any]] = []
    target_rows: list[dict[str, Any]] = []
    collateral_rows: list[dict[str, Any]] = []
    for category in ("三厢车", "跑车"):
        rule = next(
            r
            for r in rules
            if r.category == category and r.logical_size == "3L" and r.lane == "普通"
        )
        result = evaluate_candidate(rule, 4693, rules, base_rows, base_metrics)
        improved = set(result["improved_families"])
        worsened = set(result["worsened_families"])
        for before, after in result["changed"]:
            key = family_key(before)
            is_accord = clean_text(before.get("MAKE")) == "Honda" and clean_text(before.get("MODEL")) == "Accord"
            impact_type = (
                "目标：Accord 跨代归组"
                if is_accord
                else "连带但改善该车型跨代"
                if key in improved
                else "跨代纯度下降"
                if key in worsened
                else "纯连带换档"
            )
            item = {
                "影响类型": impact_type,
                "MAKE": clean_text(before.get("MAKE")),
                "MODEL": clean_text(before.get("MODEL")),
                "结构": clean_text(before.get("结构")),
                "代际": clean_text(before.get("代际")),
                "YEAR": clean_text(before.get("YEAR")),
                "分类": category,
                "L-MM": before["__length"],
                "参考插片": before["__rough"],
                "销量合计": before["__sales"],
                "当前逻辑尺码": before["__logical_size"],
                "模拟逻辑尺码": after["__logical_size"],
                "当前长度余量-mm": before["__length_surplus"],
                "模拟长度余量-mm": after["__length_surplus"],
                "余量增加-mm": (
                    after["__length_surplus"] - before["__length_surplus"]
                    if after["__length_surplus"] is not None and before["__length_surplus"] is not None
                    else None
                ),
                "边界方案": "3L 长上限 4736 → 4693 mm",
                "建议处理": "手动改为 3XL" if is_accord else "保持当前自动尺码，不随全局规则换档",
                "DIMENSION-ID": clean_text(before.get("DIMENSION-ID")),
                "源行号": before["__source_row"],
            }
            all_changed.append(item)
            (target_rows if is_accord else collateral_rows).append(item)
    all_changed.sort(key=lambda row: (row["影响类型"] != "目标：Accord 跨代归组", -row["销量合计"], row["MAKE"], row["MODEL"]))
    target_rows.sort(key=lambda row: (row["分类"], row["L-MM"], row["YEAR"]))
    collateral_rows.sort(key=lambda row: (-row["销量合计"], row["MAKE"], row["MODEL"], row["YEAR"]))
    return all_changed, target_rows, collateral_rows


def w_boundary_observations(base_rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    rows = [
        row
        for row in base_rows
        if row["__category"] in {"三厢车", "跑车"}
        and row["__rough"] is not None
        and 115 <= row["__rough"] <= 124
    ]
    output = []
    for row in rows:
        output.append(
            {
                "相对120": row["__rough"] - 120,
                "MAKE": clean_text(row.get("MAKE")),
                "MODEL": clean_text(row.get("MODEL")),
                "结构": clean_text(row.get("结构")),
                "代际": clean_text(row.get("代际")),
                "YEAR": clean_text(row.get("YEAR")),
                "分类": row["__category"],
                "L-MM": row["__length"],
                "参考插片": row["__rough"],
                "销量合计": row["__sales"],
                "当前逻辑尺码": row["__logical_size"],
                "当前内部尺码": row["__internal_size"],
                "观察": "W侧" if row["__rough"] > 120 else "普通侧",
                "结论": "保持 120；边界两侧均有高销量车型，调整会直接改变普通/W版型判断",
                "DIMENSION-ID": clean_text(row.get("DIMENSION-ID")),
            }
        )
    output.sort(key=lambda row: (row["参考插片"], -row["销量合计"], row["MAKE"], row["MODEL"]))
    return output


def auto_widths(ws, minimum: int = 8, maximum: int = 36) -> None:
    for column_cells in ws.columns:
        letter = get_column_letter(column_cells[0].column)
        width = minimum
        for cell in column_cells:
            if cell.value is None:
                continue
            text = str(cell.value)
            width = max(width, min(maximum, len(text) * (1.6 if re.search(r"[\u4e00-\u9fff]", text) else 1.05) + 2))
        ws.column_dimensions[letter].width = width


def style_title(ws, title: str, subtitle: str, end_col: int) -> None:
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=end_col)
    ws["A1"] = title
    ws["A1"].font = Font(name="Microsoft YaHei", size=18, bold=True, color=WHITE)
    ws["A1"].fill = PatternFill("solid", fgColor=NAVY)
    ws["A1"].alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 32
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=end_col)
    ws["A2"] = subtitle
    ws["A2"].font = Font(name="Microsoft YaHei", size=10, color="475569")
    ws["A2"].fill = PatternFill("solid", fgColor=LIGHT_GRAY)
    ws["A2"].alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    ws.row_dimensions[2].height = 28
    ws.sheet_view.showGridLines = False


def write_table_sheet(
    workbook: Workbook,
    sheet_name: str,
    title: str,
    subtitle: str,
    headers: list[str],
    rows: list[dict[str, Any] | list[Any]],
    table_name: str,
    widths: dict[str, float] | None = None,
    number_formats: dict[str, str] | None = None,
) -> Any:
    ws = workbook.create_sheet(sheet_name)
    style_title(ws, title, subtitle, len(headers))
    header_row = 4
    for col, header in enumerate(headers, start=1):
        cell = ws.cell(header_row, col, header)
        cell.font = Font(name="Microsoft YaHei", size=10, bold=True, color=WHITE)
        cell.fill = PatternFill("solid", fgColor=TEAL)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row_idx, row in enumerate(rows, start=header_row + 1):
        values = [row.get(header) for header in headers] if isinstance(row, dict) else row
        for col_idx, value in enumerate(values, start=1):
            cell = ws.cell(row_idx, col_idx, value)
            cell.font = Font(name="Microsoft YaHei", size=9, color="1E293B")
            cell.alignment = Alignment(vertical="center", wrap_text=col_idx in {len(headers)})
            if row_idx % 2 == 0:
                cell.fill = PatternFill("solid", fgColor="F8FAFC")
    end_row = max(header_row + 1, header_row + len(rows))
    if not rows:
        for col_idx in range(1, len(headers) + 1):
            ws.cell(header_row + 1, col_idx, "")
    table = Table(displayName=table_name, ref=f"A{header_row}:{get_column_letter(len(headers))}{end_row}")
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False, showRowStripes=True, showColumnStripes=False
    )
    ws.add_table(table)
    ws.freeze_panes = "A5"
    ws.row_dimensions[header_row].height = 32
    if number_formats:
        header_to_col = {header: index + 1 for index, header in enumerate(headers)}
        for header, fmt in number_formats.items():
            if header not in header_to_col:
                continue
            for row_idx in range(header_row + 1, end_row + 1):
                ws.cell(row_idx, header_to_col[header]).number_format = fmt
    auto_widths(ws)
    if widths:
        for letter, width in widths.items():
            ws.column_dimensions[letter].width = width
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.outlinePr.summaryBelow = True
    return ws


def add_decision_colors(ws, decision_col: int, start_row: int, end_row: int) -> None:
    colors = {
        "建议小范围试行": (LIGHT_TEAL, GREEN),
        "样车验证后可选": (LIGHT_ORANGE, ORANGE),
        "保持规则，车型手动": (LIGHT_BLUE, BLUE),
        "保持现状": (LIGHT_GRAY, "475569"),
    }
    for row in range(start_row, end_row + 1):
        cell = ws.cell(row, decision_col)
        if cell.value in colors:
            fill, font = colors[cell.value]
            cell.fill = PatternFill("solid", fgColor=fill)
            cell.font = Font(name="Microsoft YaHei", size=9, bold=True, color=font)


def create_summary_sheet(
    workbook: Workbook,
    source_meta: dict[str, Any],
    records: list[dict[str, Any]],
    base_rows: list[dict[str, Any]],
    mismatch_count: int,
    manual_rows: list[dict[str, Any]],
    collateral_rows: list[dict[str, Any]],
) -> Any:
    ws = workbook.active
    ws.title = "结论摘要"
    style_title(
        ws,
        "车型尺码上限微调分析",
        "结论先行：乘用车 3L 与普通/W 分界不建议全局改；Accord 等边界车型用手动清单处理。仅 PK-XXL +11 mm 具备低连带的小范围试行条件。",
        10,
    )
    ws["A4"] = "源文件"
    ws["B4"] = source_meta["path"]
    ws["A5"] = "修改时间"
    ws["B5"] = source_meta["mtime"]
    ws["A6"] = "SHA-256"
    ws["B6"] = source_meta["sha256"]
    ws.merge_cells("B4:J4")
    ws.merge_cells("B5:J5")
    ws.merge_cells("B6:J6")
    for row in range(4, 7):
        ws.cell(row, 1).font = Font(name="Microsoft YaHei", size=10, bold=True, color=NAVY)
        ws.cell(row, 2).font = Font(name="Consolas" if row == 6 else "Microsoft YaHei", size=9, color="334155")
        ws.cell(row, 2).alignment = Alignment(wrap_text=True)

    kpis = [
        ("A8", "总车型记录", len(records), "#,##0"),
        ("C8", "有分类记录", sum(1 for r in records if r["__category"]), "#,##0"),
        ("E8", "规则复算差异", mismatch_count, "#,##0"),
        ("G8", "手动候选", len(manual_rows), "#,##0"),
        ("I8", "Accord方案连带", len(collateral_rows), "#,##0"),
    ]
    for anchor, label, value, fmt in kpis:
        col = ws[anchor].column
        row = ws[anchor].row
        ws.merge_cells(start_row=row, start_column=col, end_row=row, end_column=col + 1)
        ws.cell(row, col, label)
        ws.cell(row, col).fill = PatternFill("solid", fgColor=LIGHT_GRAY)
        ws.cell(row, col).font = Font(name="Microsoft YaHei", size=9, bold=True, color="475569")
        ws.merge_cells(start_row=row + 1, start_column=col, end_row=row + 2, end_column=col + 1)
        ws.cell(row + 1, col, value)
        ws.cell(row + 1, col).number_format = fmt
        ws.cell(row + 1, col).fill = PatternFill("solid", fgColor=LIGHT_TEAL if value == 0 else LIGHT_BLUE)
        ws.cell(row + 1, col).font = Font(name="Microsoft YaHei", size=18, bold=True, color=NAVY)
        ws.cell(row + 1, col).alignment = Alignment(horizontal="center", vertical="center")

    ws["A12"] = "推荐动作"
    ws["A12"].font = Font(name="Microsoft YaHei", size=13, bold=True, color=NAVY)
    action_headers = ["优先级", "对象", "建议", "原因 / 风险控制"]
    action_rows = [
        ["1", "PK-XXL 长上限", "6451 → 6462 mm，小范围试行", "只影响 2 条 F-250/F-350 Super Duty 记录；先确认向小一档后仍能覆盖。"],
        ["2", "Honda Accord", "不改 3L 全局上限；4 条边界代际手动改到 3XL", "全局 4736 → 4693 会连带 65 条非 Accord 记录换大一档。"],
        ["3", "普通/W 分界", "维持参考插片上限 120", "±5 区间已有 18 条、330.6 万销量，且两侧车型几何形态差异明显。"],
        ["4", "其余长上限", "保持现状，使用手动清单逐车型处理", "多数候选会移动几十条记录，或把长车分到更小尺码，需实物验证。"],
    ]
    start = 13
    for col, header in enumerate(action_headers, 1):
        cell = ws.cell(start, col, header)
        cell.fill = PatternFill("solid", fgColor=TEAL)
        cell.font = Font(name="Microsoft YaHei", size=10, bold=True, color=WHITE)
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for r_idx, row in enumerate(action_rows, start + 1):
        for c_idx, value in enumerate(row, 1):
            cell = ws.cell(r_idx, c_idx, value)
            cell.font = Font(name="Microsoft YaHei", size=10, color="1E293B")
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if r_idx % 2 == 0:
                cell.fill = PatternFill("solid", fgColor="F8FAFC")
    ws.merge_cells(start_row=start, start_column=4, end_row=start, end_column=10)
    for r_idx in range(start + 1, start + 1 + len(action_rows)):
        ws.merge_cells(start_row=r_idx, start_column=4, end_row=r_idx, end_column=10)
        ws.row_dimensions[r_idx].height = 34
    ws.column_dimensions["A"].width = 12
    ws.column_dimensions["B"].width = 24
    ws.column_dimensions["C"].width = 38
    for letter in "DEFGHIJ":
        ws.column_dimensions[letter].width = 14
    ws.freeze_panes = "A4"
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    return ws


def create_analysis_workbook(
    rules: list[SizeRule],
    records: list[dict[str, Any]],
    rule_snapshot: list[list[Any]],
    best_results: list[dict[str, Any]],
    manual_rows: list[dict[str, Any]],
    accord_collateral: list[dict[str, Any]],
    source_meta: dict[str, Any],
    mismatch_count: int,
) -> None:
    base_rows = assign_records(records, rules)
    workbook = Workbook()
    create_summary_sheet(workbook, source_meta, records, base_rows, mismatch_count, manual_rows, accord_collateral)

    scenario_overrides: dict[tuple[str, str, str], dict[str, Any]] = {}
    base_metrics = family_metrics(base_rows, "__logical_size")
    for category in ("三厢车", "跑车"):
        rule = next(r for r in rules if r.category == category and r.logical_size == "3L" and r.lane == "普通")
        scenario_overrides[(category, "普通", "3L")] = evaluate_candidate(rule, 4693, rules, base_rows, base_metrics)

    recommendation_rows = []
    for result in best_results:
        rule = result["rule"]
        result = scenario_overrides.get((rule.category, rule.lane, rule.logical_size), result)
        decision, effective_upper, reason = policy_decision(result)
        recommendation_rows.append(
            {
                "源尺码行": rule.row_no,
                "分类": rule.category,
                "通道": rule.lane,
                "逻辑尺码": rule.logical_size,
                "内部尺码": rule.internal_size,
                "现行长上限": rule.length_upper,
                "模拟候选上限": result["new_upper"],
                "模拟Δ-mm": None,
                "建议落地上限": effective_upper,
                "决策": decision,
                "换档记录数": result["changed_count"],
                "换档销量": result["moved_sales"],
                "跨代归组改善销量": result["improve_sales"],
                "跨代纯度下降销量": result["worsen_sales"],
                "非目标连带销量": result["neutral_sales"],
                "±25mm边界记录": result["near_boundary_count"],
                "±25mm边界销量": result["near_boundary_sales"],
                "改善车型": "; ".join(family_label(key) for key in result["improved_families"][:8]),
                "结论说明": reason,
            }
        )
    rec_headers = list(recommendation_rows[0].keys())
    ws = write_table_sheet(
        workbook,
        "上限建议",
        "各尺码长上限：最佳模拟与保守落地建议",
        "模拟范围：乘用车 ±100 mm、皮卡 ±150 mm；销量改善衡量高销量家族的主组覆盖增加，非目标连带销量用于抑制过度换档。",
        rec_headers,
        recommendation_rows,
        "BoundaryRecommendations",
        widths={"R": 42, "S": 52},
        number_formats={
            "现行长上限": "0",
            "模拟候选上限": "0",
            "模拟Δ-mm": "+0;-0;0",
            "建议落地上限": "0",
            "换档记录数": "#,##0",
            "换档销量": "#,##0",
            "跨代归组改善销量": "#,##0",
            "跨代纯度下降销量": "#,##0",
            "非目标连带销量": "#,##0",
            "±25mm边界记录": "#,##0",
            "±25mm边界销量": "#,##0",
        },
    )
    delta_col = rec_headers.index("模拟Δ-mm") + 1
    current_col = rec_headers.index("现行长上限") + 1
    simulated_col = rec_headers.index("模拟候选上限") + 1
    for row in range(5, 5 + len(recommendation_rows)):
        ws.cell(row, delta_col, f"={get_column_letter(simulated_col)}{row}-{get_column_letter(current_col)}{row}")
        ws.cell(row, delta_col).number_format = "+0;-0;0"
    add_decision_colors(ws, rec_headers.index("决策") + 1, 5, 4 + len(recommendation_rows))

    split_rows = []
    for item in top_split_families(base_rows)[:80]:
        split_rows.append(
            {
                "车型家族": item["family"],
                "分类": item["category"],
                "通道": item["lane"],
                "代际数": item["generations"],
                "家族总销量": item["total_sales"],
                "现有尺码分布(销量)": item["sizes"],
                "销量主组": item["modal_size"],
                "主组占比": item["purity"],
                "处理原则": "边界代际手动归主组" if item["purity"] >= 0.6 else "保留多尺码；代际尺寸差异较大",
            }
        )
    write_table_sheet(
        workbook,
        "跨代分组",
        "高销量车型跨代尺码分布",
        "按 MAKE + MODEL + 结构 + 分类 + 普通/W通道聚合；同名车型跨越多代但尺寸跨度过大时，不强制并组。",
        list(split_rows[0].keys()),
        split_rows,
        "FamilyClusters",
        widths={"A": 38, "F": 42, "I": 34},
        number_formats={"家族总销量": "#,##0", "主组占比": "0.0%"},
    )

    proposed_by_row = {35: 6462}
    snapshot_rows = []
    for row_no, row in enumerate(rule_snapshot[1:], start=2):
        if not any(value is not None for value in row):
            continue
        snapshot_rows.append(
            {
                "源行号": row_no,
                "逻辑尺码": row[0],
                "内部尺码": row[1],
                "档位序号": row[2],
                "分类": row[3],
                "CAB": row[4],
                "版本": row[5],
                "现行长上限": row[6],
                "参考插片上限": row[7],
                "使用": row[8],
                "建议长上限": proposed_by_row.get(row_no, row[6]),
                "建议": "小范围试行" if row_no in proposed_by_row else "保持",
            }
        )
    write_table_sheet(
        workbook,
        "规则快照",
        "“尺码”工作表 A:I 规则快照",
        "源表不做修改；建议列仅记录本次评估结论。参考插片普通/W分界 120 保持不变。",
        list(snapshot_rows[0].keys()),
        snapshot_rows,
        "RuleSnapshot",
        widths={"B": 14, "C": 16, "E": 12, "L": 18},
        number_formats={"现行长上限": "0", "参考插片上限": "0", "建议长上限": "0"},
    )

    category_rows = []
    for category in ["两厢车", "跑车", "三厢车", "越野车", "皮卡", "未分类"]:
        key = "" if category == "未分类" else category
        rows = [row for row in base_rows if row["__category"] == key]
        category_rows.append(
            {
                "分类": category,
                "记录数": len(rows),
                "销量合计": sum(row["__sales"] for row in rows),
                "已匹配尺码": sum(1 for row in rows if row["__logical_size"]),
                "未匹配尺码": sum(1 for row in rows if not row["__logical_size"]),
                "未匹配销量": sum(row["__sales"] for row in rows if not row["__logical_size"]),
            }
        )
    write_table_sheet(
        workbook,
        "数据质量",
        "分析覆盖与规则复算质量",
        f"现行规则复算与“自动尺码”字段差异：{mismatch_count} 条。未分类的 450 条记录销量均为 0，不进入上限优化。",
        list(category_rows[0].keys()),
        category_rows,
        "DataQuality",
        number_formats={"记录数": "#,##0", "销量合计": "#,##0", "已匹配尺码": "#,##0", "未匹配尺码": "#,##0", "未匹配销量": "#,##0"},
    )

    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    workbook.calculation.calcMode = "auto"
    workbook.save(ANALYSIS_BOOK)


def create_manual_workbook(
    manual_rows: list[dict[str, Any]],
    accord_targets: list[dict[str, Any]],
    accord_collateral: list[dict[str, Any]],
    w_rows: list[dict[str, Any]],
    source_meta: dict[str, Any],
) -> None:
    workbook = Workbook()
    ws = workbook.active
    ws.title = "使用说明"
    style_title(
        ws,
        "边界负面影响及手动调整清单",
        "此工作簿不修改源规则。先处理 P0 Accord 专项，再按销量与风险逐项核对 P1/P2；“验证后手动降一档”必须先做样车验证。",
        8,
    )
    notes = [
        ("A4", "源文件", source_meta["path"]),
        ("A5", "源修改时间", source_meta["mtime"]),
        ("A6", "人工执行顺序", "1) Accord专项；2) 手动调整清单 P1；3) 查看边界负面影响，确认不做全局 3L 下调；4) W版边界观察仅供复核。"),
        ("A8", "风险口径", "升一档：覆盖更宽松但可能偏松；降一档：可能偏紧，必须样车验证。销量用于排序，不代表实物适配证据。"),
    ]
    for anchor, label, value in notes:
        cell = ws[anchor]
        cell.value = label
        cell.font = Font(name="Microsoft YaHei", size=10, bold=True, color=NAVY)
        value_cell = ws.cell(cell.row, 2, value)
        ws.merge_cells(start_row=cell.row, start_column=2, end_row=cell.row, end_column=8)
        value_cell.font = Font(name="Microsoft YaHei", size=10, color="334155")
        value_cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[cell.row].height = 32 if cell.row in {6, 8} else 22
    ws.column_dimensions["A"].width = 18
    for letter in "BCDEFGH":
        ws.column_dimensions[letter].width = 16
    ws.sheet_view.showGridLines = False

    manual_headers = list(manual_rows[0].keys()) if manual_rows else []
    mws = write_table_sheet(
        workbook,
        "手动调整清单",
        "高销量跨代车型：边界代际手动调整候选",
        "仅列相邻尺码、距边界较近的代际。P0 为 Accord；P1 为单行销量 ≥25 万；降一档项目必须先验证。",
        manual_headers,
        manual_rows,
        "ManualAdjustments",
        widths={"X": 56, "Y": 48},
        number_formats={
            "L-MM": "0",
            "参考插片": "0",
            "销量合计": "#,##0",
            "距当前边界-mm": "0",
            "建议后长度余量-mm": "0",
            "家族总销量": "#,##0",
            "当前主组占比": "0.0%",
            "源行号": "0",
        },
    )
    for row in range(5, 5 + len(manual_rows)):
        priority = mws.cell(row, 1).value
        fill = LIGHT_RED if priority == "P0" else LIGHT_ORANGE if priority == "P1" else LIGHT_BLUE
        mws.cell(row, 1).fill = PatternFill("solid", fgColor=fill)
        mws.cell(row, 1).font = Font(name="Microsoft YaHei", size=9, bold=True, color=NAVY)

    accord_headers = list(accord_targets[0].keys()) if accord_targets else []
    aws = write_table_sheet(
        workbook,
        "Accord专项",
        "Honda Accord 跨代归入 3XL：建议手动处理",
        "建议只处理接近 3L 上边界的 4 条记录；更早、更短的代际继续保留 3M/3L，避免 3XL 余量过大。",
        accord_headers,
        accord_targets,
        "AccordTargets",
        widths={"Q": 32, "R": 52},
        number_formats={"L-MM": "0", "参考插片": "0", "销量合计": "#,##0", "当前长度余量-mm": "0", "模拟长度余量-mm": "0", "余量增加-mm": "0", "源行号": "0"},
    )
    for row in range(5, 5 + len(accord_targets)):
        aws.cell(row, 1).fill = PatternFill("solid", fgColor=LIGHT_TEAL)

    collateral_headers = list(accord_collateral[0].keys()) if accord_collateral else []
    cws = write_table_sheet(
        workbook,
        "边界负面影响",
        "若全局把 3L 上限 4736 → 4693 mm：非 Accord 连带影响",
        "这些记录会被迫从 3L 换到 3XL；建议保持其当前自动尺码。按销量从高到低排列，便于人工核查。",
        collateral_headers,
        accord_collateral,
        "BoundaryCollateral",
        widths={"A": 22, "Q": 38, "R": 50},
        number_formats={"L-MM": "0", "参考插片": "0", "销量合计": "#,##0", "当前长度余量-mm": "0", "模拟长度余量-mm": "0", "余量增加-mm": "0", "源行号": "0"},
    )
    for row in range(5, 5 + len(accord_collateral)):
        impact = cws.cell(row, 1).value
        if impact == "跨代纯度下降":
            cws.cell(row, 1).fill = PatternFill("solid", fgColor=LIGHT_RED)
        elif impact == "纯连带换档":
            cws.cell(row, 1).fill = PatternFill("solid", fgColor=LIGHT_ORANGE)
        else:
            cws.cell(row, 1).fill = PatternFill("solid", fgColor=LIGHT_BLUE)

    w_headers = list(w_rows[0].keys()) if w_rows else []
    write_table_sheet(
        workbook,
        "W版边界观察",
        "参考插片 120 附近车型",
        "列出 115–124 区间。120 两侧同时存在高销量车型；将上限改到 124 会把 Oldsmobile Cutlass 等从 W 版切回普通版。",
        w_headers,
        w_rows,
        "WideBoundary",
        widths={"M": 30},
        number_formats={"相对120": "+0;-0;0", "L-MM": "0", "参考插片": "0", "销量合计": "#,##0"},
    )

    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    workbook.calculation.calcMode = "auto"
    workbook.save(MANUAL_BOOK)


def generate_deliverables(
    rules: list[SizeRule], records: list[dict[str, Any]], rule_snapshot: list[list[Any]], best: list[dict[str, Any]]
) -> dict[str, Any]:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    base_rows = assign_records(records, rules)
    mismatches = verify_current_assignments(records, rules)
    manual_rows = manual_adjustment_candidates(base_rows, rules)
    accord_all, accord_targets, accord_collateral = accord_boundary_scenario(base_rows, rules)
    w_rows = w_boundary_observations(base_rows)
    stat = SIZE_SOURCE.stat()
    rule_stat = RULE_SOURCE.stat()
    source_meta = {
        "path": str(SIZE_SOURCE.relative_to(ROOT)),
        "mtime": datetime.fromtimestamp(stat.st_mtime).strftime("%Y-%m-%d %H:%M:%S"),
        "sha256": sha256_file(SIZE_SOURCE),
        "rule_path": str(RULE_SOURCE.relative_to(ROOT)),
        "rule_mtime": datetime.fromtimestamp(rule_stat.st_mtime).strftime("%Y-%m-%d %H:%M:%S"),
        "rule_sha256": sha256_file(RULE_SOURCE),
    }
    create_analysis_workbook(
        rules,
        records,
        rule_snapshot,
        best,
        manual_rows,
        accord_collateral,
        source_meta,
        len(mismatches),
    )
    create_manual_workbook(manual_rows, accord_targets, accord_collateral, w_rows, source_meta)
    summary = {
        "source": source_meta,
        "records": len(records),
        "active_rules": len(rules),
        "mismatches": len(mismatches),
        "manual_candidates": len(manual_rows),
        "accord_targets": len(accord_targets),
        "accord_collateral": len(accord_collateral),
        "accord_collateral_sales": sum(row["销量合计"] for row in accord_collateral),
        "w_boundary_rows_115_124": len(w_rows),
        "w_boundary_sales_115_124": sum(row["销量合计"] for row in w_rows),
        "analysis_book": str(ANALYSIS_BOOK),
        "manual_book": str(MANUAL_BOOK),
    }
    print(json.dumps(summary, ensure_ascii=False, indent=2))
    return summary


def top_split_families(base_rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    metrics = family_metrics(base_rows, "__logical_size")
    output = []
    for key, metric in metrics.items():
        if metric["generations"] < 2 or len(metric["sizes"]) < 2:
            continue
        size_sales: dict[str, float] = defaultdict(float)
        for row in metric["rows"]:
            size_sales[row["__logical_size"]] += row["__sales"]
        output.append(
            {
                "key": key,
                "family": family_label(key),
                "category": key[3],
                "lane": key[4],
                "generations": metric["generations"],
                "total_sales": metric["total_sales"],
                "sizes": ", ".join(f"{size}:{int(sales):,}" for size, sales in sorted(size_sales.items())),
                "modal_size": metric["modal_size"],
                "purity": metric["purity"],
                "rows": metric["rows"],
            }
        )
    output.sort(key=lambda x: (-x["total_sales"], x["family"]))
    return output


def print_exploration(
    rules: list[SizeRule], records: list[dict[str, Any]], best: list[dict[str, Any]]
) -> None:
    base_rows = assign_records(records, rules)
    mismatches = verify_current_assignments(records, rules)
    print(
        json.dumps(
            {
                "size_source": str(SIZE_SOURCE),
                "rule_source": str(RULE_SOURCE),
                "source_rows": len(records),
                "active_rules": len(rules),
                "match_mismatches": len(mismatches),
                "category_counts": Counter(r["__category"] for r in records),
            },
            ensure_ascii=False,
            default=excel_serialized,
            indent=2,
        )
    )
    print("\nTOP SPLIT FAMILIES")
    for item in top_split_families(base_rows)[:30]:
        print(
            f"{item['family']} | {item['category']}/{item['lane']} | "
            f"sales={item['total_sales']:,.0f} | {item['sizes']} | modal={item['modal_size']} "
            f"purity={item['purity']:.1%}"
        )
    print("\nBEST BOUNDARY CANDIDATES")
    for result in best:
        rule = result["rule"]
        decision, reason = classify_candidate(result)
        improved = "; ".join(family_label(k) for k in result["improved_families"][:4])
        print(
            f"{rule.category}/{rule.lane}/{rule.logical_size} row={rule.row_no}: "
            f"{rule.length_upper:.0f}->{result['new_upper']} ({result['delta']:+.0f}) | "
            f"score={result['net_score']:,.0f} improve={result['improve_sales']:,.0f} "
            f"worsen={result['worsen_sales']:,.0f} neutral={result['neutral_sales']:,.0f} "
            f"changed={result['changed_count']} near={result['near_boundary_count']} | "
            f"{decision} | {improved} | {reason}"
        )


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Analyze vehicle-cover size boundaries.")
    parser.add_argument("--explore", action="store_true", help="Print analysis without creating deliverables.")
    parser.add_argument(
        "--export-adjustments-csv",
        action="store_true",
        help="Export all manual-adjustment candidates to the four-column CSV.",
    )
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    rules, records, rule_snapshot = read_source(RULE_SOURCE, SIZE_SOURCE)
    best, _ = analyze_candidates(rules, records)
    if args.explore:
        print_exploration(rules, records, best)
        return
    if args.export_adjustments_csv:
        base_rows = assign_records(records, rules)
        manual_rows = manual_adjustment_candidates(base_rows, rules)
        print(json.dumps(export_manual_adjustments_csv(manual_rows), ensure_ascii=False, indent=2))
        return
    generate_deliverables(rules, records, rule_snapshot, best)


if __name__ == "__main__":
    main()
