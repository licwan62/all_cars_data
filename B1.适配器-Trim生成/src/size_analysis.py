from __future__ import annotations

import json
import warnings
from collections import Counter, defaultdict
from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook

from src.trimlist import clean, read_csv, sha256_file, write_csv


warnings.filterwarnings(
    "ignore", message="Unknown extension is not supported and will be removed"
)
warnings.filterwarnings(
    "ignore",
    message="Conditional Formatting extension is not supported and will be removed",
)


NON_PUBLISHABLE_SIZES = {"无可用尺码", "数据不全"}

ANALYSIS_HEADER = [
    "Year",
    "Make",
    "Model",
    "DIMENSION-ID数量",
    "映射行数",
    "原始自动尺码数量",
    "原始自动尺码列表",
    "有效Size数量",
    "有效Size列表",
    "状态值列表",
    "源主车型数量",
    "源主车型列表",
    "结构列表",
    "尺码系列",
    "匹配方式",
    "判定",
    "展开类型",
    "发布处理",
]

MULTI_SIZE_HEADER = [
    "Year",
    "Make",
    "Model",
    "Size数量",
    "Size列表",
    "DIMENSION-ID数量",
    "DIMENSION-ID列表",
    "源主车型列表",
    "版本列表",
    "结构列表",
    "CAB列表",
    "BED列表",
    "分类列表",
    "尺码系列",
    "匹配方式",
    "审核状态",
    "展开类型",
    "发布判定",
    "说明",
]

MULTI_SIZE_DETAIL_HEADER = [
    "Year",
    "Make",
    "Model",
    "Size",
    "是否可发布Size",
    "DIMENSION-ID",
    "源MAKE",
    "源MODEL",
    "版本",
    "结构",
    "CAB",
    "BED",
    "分类",
    "匹配方式",
    "审核状态",
    "证据URL",
    "展开类型",
]

DIMENSION_SIZE_HEADER = [
    "DIMENSION-ID",
    "Size",
    "源MAKE",
    "源MODEL",
    "版本",
    "结构",
    "CAB",
    "BED",
    "YEAR",
    "分类",
]

DIMENSION_TRIM_HEADER = [
    "DIMENSION-ID",
    "Trims",
]

ADAPTER_SIZE_HEADER = [
    "DIMENSION-ID",
    "Size",
    "Year",
    "Make",
    "Model",
    "版本",
    "结构",
    "匹配方式",
    "审核状态",
    "证据URL",
]

FINAL_ADAPTER_HEADER = [
    "DIMENSION-ID",
    "Size",
    "Year",
    "Make",
    "Model",
]

NO_SIZE_HEADER = [
    "Year",
    "Make",
    "Model",
    "DIMENSION-ID数量",
    "DIMENSION-ID列表",
    "状态值列表",
    "源主车型列表",
    "匹配方式",
    "建议处理",
]

SIZE_WORKBOOK_COLUMNS = [
    "DIMENSION-ID",
    "自动尺码",
    "MAKE",
    "MODEL",
    "TRIM",
    "版本",
    "结构",
    "CAB",
    "BED",
    "YEAR",
    "分类",
]


@dataclass
class SizeAnalysisResult:
    analysis_rows: list[dict[str, object]]
    multi_size_rows: list[dict[str, object]]
    multi_size_detail_rows: list[dict[str, object]]
    no_size_rows: list[dict[str, object]]
    dimension_size_rows: list[dict[str, object]]
    dimension_trim_rows: list[dict[str, object]]
    adapter_size_rows: list[dict[str, object]]
    final_adapter_rows: list[dict[str, object]]
    report: dict[str, object]


def load_size_source(path: Path, sheet_name: str = "尺码匹配") -> list[dict[str, str]]:
    """读取统一尺码分析 CSV；旧 Excel 工作簿仍作为兼容输入。"""
    if path.suffix.lower() == ".csv":
        source_rows = read_csv(path)
        if not source_rows:
            raise ValueError(f"{path} 没有数据行")
        missing = [
            column for column in SIZE_WORKBOOK_COLUMNS if column not in source_rows[0]
        ]
        if missing:
            raise ValueError(f"{path} 缺少字段: {', '.join(missing)}")
        return [
            {column: clean(row.get(column)) for column in SIZE_WORKBOOK_COLUMNS}
            for row in source_rows
            if clean(row.get("DIMENSION-ID"))
        ]

    with warnings.catch_warnings():
        warnings.simplefilter("ignore", UserWarning)
        workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if sheet_name not in workbook.sheetnames:
            raise ValueError(f"{path} 不存在工作表: {sheet_name}")
        sheet = workbook[sheet_name]
        iterator = sheet.iter_rows(values_only=True)
        try:
            header_values = next(iterator)
        except StopIteration as exc:
            raise ValueError(f"{path} 工作表为空: {sheet_name}") from exc
        indexes = {
            clean(value): index
            for index, value in enumerate(header_values)
            if clean(value)
        }
        missing = [column for column in SIZE_WORKBOOK_COLUMNS if column not in indexes]
        if missing:
            raise ValueError(
                f"{path} {sheet_name} 缺少字段: {', '.join(missing)}"
            )

        rows: list[dict[str, str]] = []
        for values in iterator:
            dimension_id = clean(values[indexes["DIMENSION-ID"]])
            if not dimension_id:
                continue
            rows.append(
                {
                    column: clean(values[indexes[column]])
                    for column in SIZE_WORKBOOK_COLUMNS
                }
            )
        return rows
    finally:
        workbook.close()


def load_size_workbook(path: Path, sheet_name: str) -> list[dict[str, str]]:
    """兼容旧调用名；新代码应使用 load_size_source。"""
    return load_size_source(path, sheet_name)


def size_family(value: str) -> str:
    size = clean(value)
    if not size or size in NON_PUBLISHABLE_SIZES:
        return "STATUS"
    if size.startswith("PK-"):
        return "PK"
    if size.startswith("Y"):
        return "Y"
    if size.startswith("2"):
        return "2"
    if size.startswith("3"):
        return "3"
    return "OTHER"


def joined(values: set[str]) -> str:
    return "; ".join(sorted(value for value in values if value))


def normalized_trims(value: object) -> str:
    trims: list[str] = []
    seen: set[str] = set()
    for item in clean(value).replace("；", ",").replace(";", ",").split(","):
        trim = clean(item)
        key = trim.casefold()
        if not trim or key in seen:
            continue
        seen.add(key)
        trims.append(trim)
    return " | ".join(trims)


def classify_expansion(
    source_models: set[str],
    versions: set[str],
    structures: set[str],
) -> tuple[str, str]:
    flags: list[str] = []
    if len(structures) > 1:
        flags.append("MULTI_STRUCTURE")
    if len(versions) > 1:
        flags.append("MULTI_VERSION")
    if len(source_models) > 1:
        flags.append("MULTI_SOURCE_MODEL")
    if not flags:
        flags.append("MULTI_DIMENSION_VARIANT")
    return "; ".join(flags), "保留并展开全部 DIMENSION-ID + Size"


def build_size_analysis(
    trim_rows: list[dict[str, str]],
    audit_rows: list[dict[str, str]],
    size_rows: list[dict[str, str]],
) -> SizeAnalysisResult:
    size_counts = Counter(row["DIMENSION-ID"] for row in size_rows)
    duplicate_size_ids = sorted(key for key, count in size_counts.items() if count > 1)
    if duplicate_size_ids:
        raise ValueError(
            f"车型数据尺码存在重复 DIMENSION-ID: {duplicate_size_ids[:10]}"
        )
    size_map = {row["DIMENSION-ID"]: row for row in size_rows}
    dimension_trim_rows = sorted(
        (
            {
                "DIMENSION-ID": row["DIMENSION-ID"],
                "Trims": normalized_trims(row["TRIM"]),
            }
            for row in size_rows
        ),
        key=lambda row: str(row["DIMENSION-ID"]),
    )
    audit_map = {
        (
            row["DIMENSION-ID"],
            int(row["Year"]),
            row["Make"],
            row["Model"],
        ): row
        for row in audit_rows
    }

    missing_dimension_ids: set[str] = set()
    missing_audit_keys: set[tuple[str, int, str, str]] = set()
    assignments: list[dict[str, object]] = []
    for trim in trim_rows:
        dimension_id = trim["DIMENSION-ID"]
        size_row = size_map.get(dimension_id)
        if size_row is None:
            missing_dimension_ids.add(dimension_id)
            continue
        year = int(trim["Year"])
        size = size_row["自动尺码"]
        audit_key = (dimension_id, year, trim["Make"], trim["Model"])
        audit = audit_map.get(audit_key, {})
        if not audit:
            missing_audit_keys.add(audit_key)
        assignments.append(
            {
                "Year": year,
                "Make": trim["Make"],
                "Model": trim["Model"],
                "Size": size,
                "是否可发布Size": "Y"
                if size and size not in NON_PUBLISHABLE_SIZES
                else "N",
                "DIMENSION-ID": dimension_id,
                "源MAKE": size_row["MAKE"],
                "源MODEL": size_row["MODEL"],
                "版本": size_row["版本"],
                "结构": size_row["结构"],
                "CAB": size_row["CAB"],
                "BED": size_row["BED"],
                "分类": size_row["分类"],
                "匹配方式": audit.get("匹配方式", ""),
                "审核状态": audit.get("审核状态", ""),
                "证据URL": audit.get("证据URL", ""),
            }
        )

    groups: dict[tuple[int, str, str], list[dict[str, object]]] = defaultdict(list)
    for row in assignments:
        groups[(int(row["Year"]), str(row["Make"]), str(row["Model"]))].append(row)

    analysis_rows: list[dict[str, object]] = []
    multi_size_rows: list[dict[str, object]] = []
    multi_size_detail_rows: list[dict[str, object]] = []
    no_size_rows: list[dict[str, object]] = []

    multi_dimension_keys = 0
    safe_multi_dimension_same_size = 0
    status_assignment_rows = 0
    status_only_keys = 0
    mixed_status_and_size_keys = 0
    multi_size_distribution: Counter[int] = Counter()
    multi_size_type_counts: Counter[str] = Counter()
    multi_size_makes: Counter[str] = Counter()
    multi_structure_expansion_keys = 0

    for (year, make, model), rows in sorted(groups.items()):
        dimension_ids = {str(row["DIMENSION-ID"]) for row in rows}
        raw_sizes = {str(row["Size"]) for row in rows if clean(row["Size"])}
        valid_sizes = {
            str(row["Size"])
            for row in rows
            if row["是否可发布Size"] == "Y"
        }
        status_values = raw_sizes & NON_PUBLISHABLE_SIZES
        source_models = {
            f"{row['源MAKE']} {row['源MODEL']}".strip() for row in rows
        }
        versions = {str(row["版本"]) for row in rows if clean(row["版本"])}
        structures = {str(row["结构"]) for row in rows if clean(row["结构"])}
        cabs = {str(row["CAB"]) for row in rows if clean(row["CAB"])}
        beds = {str(row["BED"]) for row in rows if clean(row["BED"])}
        categories = {str(row["分类"]) for row in rows if clean(row["分类"])}
        methods = {str(row["匹配方式"]) for row in rows if clean(row["匹配方式"])}
        review_states = {
            str(row["审核状态"]) for row in rows if clean(row["审核状态"])
        }
        families = {size_family(size) for size in valid_sizes}
        status_rows = sum(row["是否可发布Size"] == "N" for row in rows)
        status_assignment_rows += status_rows

        expansion_type = ""
        action = ""
        if len(valid_sizes) > 1:
            judgement = "MULTI_SIZE_EXPANDED"
            expansion_type, action = classify_expansion(
                source_models, versions, structures
            )
        elif len(valid_sizes) == 1:
            judgement = "SINGLE_SIZE"
            action = "保留 DIMENSION-ID + Size 后发布"
        else:
            judgement = "NO_PUBLISHABLE_SIZE"
            action = "阻止发布；先补齐或修正自动尺码"

        if len(dimension_ids) > 1:
            multi_dimension_keys += 1
            if len(valid_sizes) == 1 and not status_values:
                safe_multi_dimension_same_size += 1
        if not valid_sizes:
            status_only_keys += 1
        elif status_values:
            mixed_status_and_size_keys += 1

        analysis_rows.append(
            {
                "Year": year,
                "Make": make,
                "Model": model,
                "DIMENSION-ID数量": len(dimension_ids),
                "映射行数": len(rows),
                "原始自动尺码数量": len(raw_sizes),
                "原始自动尺码列表": joined(raw_sizes),
                "有效Size数量": len(valid_sizes),
                "有效Size列表": joined(valid_sizes),
                "状态值列表": joined(status_values),
                "源主车型数量": len(source_models),
                "源主车型列表": joined(source_models),
                "结构列表": joined(structures),
                "尺码系列": joined(families),
                "匹配方式": joined(methods),
                "判定": judgement,
                "展开类型": expansion_type,
                "发布处理": action,
            }
        )

        if judgement == "MULTI_SIZE_EXPANDED":
            multi_size_distribution[len(valid_sizes)] += 1
            multi_size_makes[make] += 1
            flags = set(expansion_type.split("; "))
            for flag in flags:
                if flag:
                    multi_size_type_counts[flag] += 1
            if "MULTI_STRUCTURE" in flags:
                multi_structure_expansion_keys += 1

            multi_size_rows.append(
                {
                    "Year": year,
                    "Make": make,
                    "Model": model,
                    "Size数量": len(valid_sizes),
                    "Size列表": joined(valid_sizes),
                    "DIMENSION-ID数量": len(dimension_ids),
                    "DIMENSION-ID列表": joined(dimension_ids),
                    "源主车型列表": joined(source_models),
                    "版本列表": joined(versions),
                    "结构列表": joined(structures),
                    "CAB列表": joined(cabs),
                    "BED列表": joined(beds),
                    "分类列表": joined(categories),
                    "尺码系列": joined(families),
                    "匹配方式": joined(methods),
                    "审核状态": joined(review_states),
                    "展开类型": expansion_type,
                    "发布判定": "保留全部分支",
                    "说明": "多 Size 合法；每个 DIMENSION-ID + Size 独立保留",
                }
            )
            for row in rows:
                multi_size_detail_rows.append(
                    {
                        **row,
                        "展开类型": expansion_type,
                    }
                )

        elif judgement == "NO_PUBLISHABLE_SIZE":
            no_size_rows.append(
                {
                    "Year": year,
                    "Make": make,
                    "Model": model,
                    "DIMENSION-ID数量": len(dimension_ids),
                    "DIMENSION-ID列表": joined(dimension_ids),
                    "状态值列表": joined(status_values),
                    "源主车型列表": joined(source_models),
                    "匹配方式": joined(methods),
                    "建议处理": action,
                }
            )

    valid_assignments: list[dict[str, object]] = [
        row for row in assignments if row["是否可发布Size"] == "Y"
    ]
    final_adapter_rows = [
        {column: row[column] for column in FINAL_ADAPTER_HEADER}
        for row in assignments
    ]
    dimension_size_index: dict[tuple[str, str], dict[str, object]] = {}
    for row in size_rows:
        size = row["自动尺码"]
        if not size or size in NON_PUBLISHABLE_SIZES:
            continue
        pair = (row["DIMENSION-ID"], size)
        dimension_size_index[pair] = {
            "DIMENSION-ID": row["DIMENSION-ID"],
            "Size": size,
            "源MAKE": row["MAKE"],
            "源MODEL": row["MODEL"],
            "版本": row["版本"],
            "结构": row["结构"],
            "CAB": row["CAB"],
            "BED": row["BED"],
            "YEAR": row["YEAR"],
            "分类": row["分类"],
        }
    dimension_size_rows = sorted(
        dimension_size_index.values(),
        key=lambda row: (str(row["DIMENSION-ID"]), str(row["Size"])),
    )

    adapter_index: dict[tuple[str, str, int, str, str], dict[str, object]] = {}
    for row in valid_assignments:
        atom = (
            str(row["DIMENSION-ID"]),
            str(row["Size"]),
            int(row["Year"]),
            str(row["Make"]),
            str(row["Model"]),
        )
        adapter_index[atom] = {
            column: row[column]
            for column in ADAPTER_SIZE_HEADER
        }
    adapter_size_rows = sorted(
        adapter_index.values(),
        key=lambda row: (
            str(row["DIMENSION-ID"]),
            str(row["Size"]),
            int(row["Year"]),
            str(row["Make"]).casefold(),
            str(row["Model"]).casefold(),
        ),
    )

    multi_size_keys = len(multi_size_rows)
    unique_keys_with_valid_size = len(groups) - len(no_size_rows)

    hard_errors: list[str] = []
    if missing_dimension_ids:
        hard_errors.append(
            f"TrimList 有 {len(missing_dimension_ids)} 个 DIMENSION-ID 无法关联尺码表"
        )
    if missing_audit_keys:
        hard_errors.append(
            f"TrimList 有 {len(missing_audit_keys)} 条记录缺少审核明细"
        )
    blank_sizes = sum(not clean(row["Size"]) for row in assignments)
    if blank_sizes:
        hard_errors.append(f"关联后有 {blank_sizes} 行自动尺码为空")

    report: dict[str, object] = {
        "schema_version": "2.2",
        "counts": {
            "trim_assignment_rows": len(trim_rows),
            "joined_assignment_rows": len(assignments),
            "unique_year_make_model_keys": len(groups),
            "multi_dimension_keys": multi_dimension_keys,
            "safe_multi_dimension_same_size_keys": safe_multi_dimension_same_size,
            "status_assignment_rows": status_assignment_rows,
            "status_only_no_publishable_size_keys": status_only_keys,
            "mixed_status_and_publishable_size_keys": mixed_status_and_size_keys,
            "publishable_assignment_rows": len(valid_assignments),
            "dimension_id_size_rows": len(dimension_size_rows),
            "dimension_id_trim_rows": len(dimension_trim_rows),
            "dimension_id_trim_blank_rows": sum(
                not clean(row["Trims"]) for row in dimension_trim_rows
            ),
            "adapter_size_rows": len(adapter_size_rows),
            "final_adapter_rows": len(final_adapter_rows),
            "final_adapter_status_rows_retained": len(assignments)
            - len(valid_assignments),
            "adapter_duplicate_rows_removed": len(valid_assignments)
            - len(adapter_size_rows),
            "unique_keys_with_publishable_size": unique_keys_with_valid_size,
            "single_size_keys": unique_keys_with_valid_size - multi_size_keys,
            "multi_size_expansion_keys": multi_size_keys,
            "multi_structure_expansion_keys": multi_structure_expansion_keys,
            "multi_size_detail_rows": len(multi_size_detail_rows),
        },
        "status_values": {
            value: sum(row["Size"] == value for row in assignments)
            for value in sorted(NON_PUBLISHABLE_SIZES)
        },
        "multi_size_expansions": {
            "pct_of_keys_with_publishable_size": round(
                multi_size_keys / unique_keys_with_valid_size * 100, 2
            )
            if unique_keys_with_valid_size
            else 0,
            "size_count_distribution": {
                str(key): value
                for key, value in sorted(multi_size_distribution.items())
            },
            "expansion_type_counts": dict(sorted(multi_size_type_counts.items())),
            "top_makes": [
                {"Make": make, "multi_size_keys": count}
                for make, count in multi_size_makes.most_common(15)
            ],
        },
        "checks": {
            "size_dimension_ids_unique": not duplicate_size_ids,
            "all_trim_dimension_ids_joined": not missing_dimension_ids,
            "all_joined_auto_sizes_nonblank": blank_sizes == 0,
            "analysis_partition_complete": len(analysis_rows) == len(groups),
            "dimension_id_size_primary_key_unique": len(dimension_size_rows)
            == len(dimension_size_index),
            "dimension_id_trim_matches_size_source": len(dimension_trim_rows)
            == len(size_rows),
            "dimension_id_trim_primary_key_unique": len(
                {str(row["DIMENSION-ID"]) for row in dimension_trim_rows}
            )
            == len(dimension_trim_rows),
            "adapter_atom_primary_key_unique": len(adapter_size_rows)
            == len(adapter_index),
            "final_adapter_matches_trim_rows": len(final_adapter_rows)
            == len(trim_rows),
            "final_adapter_all_sizes_nonblank": all(
                clean(row["Size"]) for row in final_adapter_rows
            ),
            "all_published_assignments_reviewed": all(
                row["审核状态"] in {"现有精确键", "联网证据批准"}
                for row in assignments
            ),
        },
        "hard_errors": hard_errors,
    }

    multi_size_detail_rows.sort(
        key=lambda row: (
            int(row["Year"]),
            str(row["Make"]).casefold(),
            str(row["Model"]).casefold(),
            str(row["Size"]),
            str(row["DIMENSION-ID"]),
        )
    )

    return SizeAnalysisResult(
        analysis_rows=analysis_rows,
        multi_size_rows=multi_size_rows,
        multi_size_detail_rows=multi_size_detail_rows,
        no_size_rows=no_size_rows,
        dimension_size_rows=dimension_size_rows,
        dimension_trim_rows=dimension_trim_rows,
        adapter_size_rows=adapter_size_rows,
        final_adapter_rows=final_adapter_rows,
        report=report,
    )


def summary_markdown(report: dict[str, object]) -> str:
    counts = report["counts"]
    expansions = report["multi_size_expansions"]
    statuses = report["status_values"]
    return f"""# TrimList 与自动尺码关联分析

## 结论

- Trim 映射行：{counts['trim_assignment_rows']:,}
- 关联成功行：{counts['joined_assignment_rows']:,}
- 唯一 `Year + Make + Model`：{counts['unique_year_make_model_keys']:,}
- 单 Size 键：{counts['single_size_keys']:,}
- 多 Size 正常展开键：{counts['multi_size_expansion_keys']:,}
- 其中多结构展开键：{counts['multi_structure_expansion_keys']:,}
- 无可发布 Size 键：{counts['status_only_no_publishable_size_keys']:,}

多 Size 占有效尺码车型键的 **{expansions['pct_of_keys_with_publishable_size']:.2f}%**；
多 Size 本身不是冲突，全部按 `DIMENSION-ID + Size` 保留。

## 唯一性与展开

| 指标 | 数量 |
|---|---:|
| 过滤状态值后可发布行 | {counts['publishable_assignment_rows']:,} |
| 最终适配器回填行 | {counts['final_adapter_rows']:,} |
| 最终适配器保留状态行 | {counts['final_adapter_status_rows_retained']:,} |
| 唯一 DIMENSION-ID + Size | {counts['dimension_id_size_rows']:,} |
| DIMENSION-ID + Trims 映射 | {counts['dimension_id_trim_rows']:,} |
| Trims 为空 | {counts['dimension_id_trim_blank_rows']:,} |
| 完整适配原子行 | {counts['adapter_size_rows']:,} |
| 可删除完整重复行 | {counts['adapter_duplicate_rows_removed']:,} |
| 多 Size 正常展开 | {counts['multi_size_expansion_keys']:,} |

## 非尺码状态

| 状态 | 映射行 |
|---|---:|
| 无可用尺码 | {statuses.get('无可用尺码', 0):,} |
| 数据不全 | {statuses.get('数据不全', 0):,} |

这些状态不应作为正式 `Size` 发布。

## 发布规则

1. `适配器.csv` 对 `TrimList.csv` 全量回填，不丢弃 `无可用尺码` 和 `数据不全` 状态行。
2. 可发布尺码分析继续排除 `无可用尺码` 和 `数据不全`。
3. `DimensionSizeMap.csv` 以 `DIMENSION-ID + Size` 为唯一键。
4. `DimensionTrimMap.csv` 每个 `DIMENSION-ID` 一行，`Trims` 仅来自尺码分析的 `TRIM` 列。
5. 多结构或多版本导致多个 Size 时，保留并展开全部已核实分支。
6. 完整发布原子键为 `DIMENSION-ID + Size + Year + Make + Model`。
7. 无有效 Size 进入 `NoPublishableSizeReport.csv`。
"""


def build_size_analysis_files(
    trim_path: Path,
    audit_path: Path,
    size_source_path: Path,
    size_sheet: str,
    output_dir: Path,
    data_dir: Path | None = None,
) -> SizeAnalysisResult:
    data_dir = data_dir or output_dir
    trim_rows = read_csv(trim_path)
    audit_rows = read_csv(audit_path)
    size_rows = load_size_source(size_source_path, size_sheet)
    result = build_size_analysis(trim_rows, audit_rows, size_rows)
    result.report["sources"] = {
        "trim_list": {"path": str(trim_path), "sha256": sha256_file(trim_path)},
        "trim_audit": {"path": str(audit_path), "sha256": sha256_file(audit_path)},
        "size_source": {
            "path": str(size_source_path),
            "format": size_source_path.suffix.lower().lstrip("."),
            "sheet": size_sheet if size_source_path.suffix.lower() != ".csv" else None,
            "sha256": sha256_file(size_source_path),
        },
    }
    if result.report["hard_errors"]:
        raise ValueError("\n".join(result.report["hard_errors"]))

    data_dir.mkdir(parents=True, exist_ok=True)
    output_dir.mkdir(parents=True, exist_ok=True)
    write_csv(data_dir / "SizeAnalysis.csv", ANALYSIS_HEADER, result.analysis_rows)
    write_csv(
        data_dir / "MultipleSizeReport.csv",
        MULTI_SIZE_HEADER,
        result.multi_size_rows,
    )
    write_csv(
        data_dir / "MultipleSizeDetail.csv",
        MULTI_SIZE_DETAIL_HEADER,
        result.multi_size_detail_rows,
    )
    write_csv(
        data_dir / "NoPublishableSizeReport.csv",
        NO_SIZE_HEADER,
        result.no_size_rows,
    )
    write_csv(
        data_dir / "DimensionSizeMap.csv",
        DIMENSION_SIZE_HEADER,
        result.dimension_size_rows,
    )
    dimension_trim_path = output_dir / "DimensionTrimMap.csv"
    write_csv(
        dimension_trim_path,
        DIMENSION_TRIM_HEADER,
        result.dimension_trim_rows,
    )
    write_csv(
        data_dir / "AdapterSizeList.csv",
        ADAPTER_SIZE_HEADER,
        result.adapter_size_rows,
    )
    final_adapter_path = output_dir / "适配器.csv"
    write_csv(
        final_adapter_path,
        FINAL_ADAPTER_HEADER,
        result.final_adapter_rows,
    )
    result.report["outputs"] = {
        "final_adapter": {
            "path": str(final_adapter_path),
            "sha256": sha256_file(final_adapter_path),
            "rows": len(result.final_adapter_rows),
        },
        "dimension_trim_map": {
            "path": str(dimension_trim_path),
            "sha256": sha256_file(dimension_trim_path),
            "rows": len(result.dimension_trim_rows),
        },
    }
    with (data_dir / "SizeAnalysisSummary.json").open(
        "w", encoding="utf-8"
    ) as file:
        json.dump(result.report, file, ensure_ascii=False, indent=2)
        file.write("\n")
    (data_dir / "SizeAnalysisSummary.md").write_text(
        summary_markdown(result.report), encoding="utf-8"
    )
    return result
